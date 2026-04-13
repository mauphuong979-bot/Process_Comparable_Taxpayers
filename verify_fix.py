import io
from docx import Document
from processor import process_docx, normalize_compare_text
from docx.oxml.ns import qn

def verify_fix():
    path = r"d:\AI\Python\GitHub\Fix Comparable Taxpayers\Demo_2025 avg Report.docx"
    with open(path, "rb") as f:
        input_data = io.BytesIO(f.read())
    
    output_data = process_docx(input_data)
    doc = Document(output_data)
    
    print("--- Verification Results ---")
    
    # 1. Check "Bulk Rejections Summary"
    target_para = None
    for i, para in enumerate(doc.paragraphs):
        if "Bulk Rejections Summary" in para.text:
            target_para = para
            target_idx = i
            break
            
    if target_para:
        print(f"Paragraph 'Bulk Rejections Summary' found at index {target_idx}")
        print(f"Style: {target_para.style.name} (Expected: Normal)")
        
        pPr = target_para._element.get_or_add_pPr()
        outlines = pPr.xpath('./w:outlineLvl')
        print(f"Outline Level tags: {len(outlines)} (Expected: 0)")
    else:
        print("Paragraph 'Bulk Rejections Summary' NOT found")

    # 2. Check for Page Break before "Step One"
    # Note: process_docx inserts a page break as a NEW paragraph before the keyword
    step_one_found = False
    for i, para in enumerate(doc.paragraphs):
        if "Step One" in para.text:
            prev_para = doc.paragraphs[i-1] if i > 0 else None
            # Check if prev_para contains a page break
            has_break = False
            if prev_para:
                for run in prev_para.runs:
                    if 'w:br' in run._element.xml and 'w:type="page"' in run._element.xml:
                        has_break = True
                        break
            print(f"Keyword 'Step One' found at index {i}. Previous para has page break: {has_break}")
            step_one_found = True
            break
    
    if not step_one_found:
        print("Keyword 'Step One' NOT found")

    # 3. Check "Arm's Length Range"
    range_found = False
    for i, para in enumerate(doc.paragraphs):
        if "Arm's Length Range" in normalize_compare_text(para.text):
            prev_para = doc.paragraphs[i-1] if i > 0 else None
            has_break = False
            if prev_para:
                for run in prev_para.runs:
                    if 'w:br' in run._element.xml and 'w:type="page"' in run._element.xml:
                        has_break = True
                        break
            print(f"Keyword 'Arm's Length Range' found at index {i}. Previous para has page break: {has_break}")
            range_found = True
            break
    
    if not range_found:
        print("Keyword 'Arm's Length Range' NOT found")

if __name__ == "__main__":
    verify_fix()
