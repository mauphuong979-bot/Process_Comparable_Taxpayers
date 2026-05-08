import streamlit as st
import io
import time
import pandas as pd
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, Mm
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook
from processor import process_docx

# =========================
# GIAO DIỆN (UI/UX)
# =========================

st.set_page_config(
    page_title="Comparable Taxpayers Processor",
    page_icon="📄",
    layout="centered",
    initial_sidebar_state="collapsed",
    menu_items={
        'Get Help': None,
        'Report a bug': None,
        'About': "# Comparable Taxpayers Processor\nProfessional document optimization."
    }
)

# Custom CSS for Premium Look
st.markdown("""
    <style>
    .stApp {
        background: linear-gradient(135deg, #f5f7fa 0%, #c3cfe2 100%);
    }

    /* Hide Streamlit branding and GitHub link */
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    
    /* Center the main container */
    .main .block-container {
        max-width: 800px;
        padding-top: 5rem;
        background: rgba(255, 255, 255, 0.9);
        border-radius: 20px;
        box-shadow: 0 10px 30px rgba(0,0,0,0.1);
        margin-top: 2rem;
        margin-bottom: 2rem;
    }

    /* Title Styling */
    h1 {
        color: #1e3a8a;
        font-family: 'Outfit', sans-serif;
        font-weight: 800;
        text-align: center;
        margin-bottom: 0.5rem;
    }
    
    .subtitle {
        color: #64748b;
        text-align: center;
        font-size: 1.1rem;
        margin-bottom: 3rem;
    }

    /* File Uploader Container */
    .stFileUploader {
        padding: 2rem;
        border: 2px dashed #3b82f6;
        border-radius: 15px;
        transition: all 0.3s ease;
    }
    
    .stFileUploader:hover {
        border-color: #1d4ed8;
        background: rgba(59, 130, 246, 0.05);
    }

    /* Button Styling */
    .stButton>button {
        width: 100%;
        background-color: #3b82f6;
        color: white;
        border-radius: 10px;
        padding: 0.75rem;
        font-weight: 600;
        border: none;
        transition: transform 0.2s ease;
    }
    
    .stButton>button:hover {
        background-color: #2563eb;
        transform: translateY(-2px);
    }

    /* Success Message */
    .stSuccess {
        background-color: #dcfce7;
        color: #166534;
        border: 1px solid #bbf7d0;
        border-radius: 10px;
    }
    </style>
""", unsafe_allow_html=True)

def word_processor_tab():
    # Content of Comparable Taxpayers Processor
    st.markdown("<h2 style='text-align: center; color: #1e3a8a;'>Comparable Taxpayers Processor</h2>", unsafe_allow_html=True)
    st.markdown("<p class='subtitle'>Professional Word document optimization and standardization</p>", unsafe_allow_html=True)

    # Settings
    auto_process = st.checkbox("Auto-process after upload", value=True, key="auto_process_check")

    # File Uploader
    uploaded_file = st.file_uploader(
        "Drag and drop Word file (.docx) here",
        type=["docx"],
        help="Only .docx format is supported",
        key="word_proc_uploader"
    )

    if uploaded_file is not None:
        file_id = f"{uploaded_file.name}_{uploaded_file.size}"
        
        # Info about selected file
        st.info(f"Selected: **{uploaded_file.name}**")

        # Start button (Manual trigger)
        start_manual = st.button("🚀 Start Processing", key="btn_process")
        
        # Logic to decide if we should run processing
        # Run if: manual click OR (auto-process is ON and this file hasn't been processed yet)
        should_run = start_manual or (auto_process and st.session_state.get('last_processed_id') != file_id)

        if should_run:
            # Update state to prevent double execution
            st.session_state['last_processed_id'] = file_id
            
            progress_bar = st.progress(0)
            status_text = st.empty()

            try:
                # UX Steps
                status_text.text("Loading document...")
                time.sleep(0.5)
                progress_bar.progress(20)

                status_text.text("Normalizing headings and fonts...")
                # Important: Read file into memory
                input_data = uploaded_file.getvalue()
                input_stream = io.BytesIO(input_data)
                progress_bar.progress(50)

                # Core processing
                processed_stream = process_docx(input_stream)
                
                status_text.text("Finalizing file...")
                progress_bar.progress(90)
                time.sleep(0.3)
                progress_bar.progress(100)
                
                # Store results in session state for persistence across reruns
                st.session_state['processed_data'] = processed_stream.getvalue()
                time_tag = datetime.now().strftime("%H%M")
                base_name = uploaded_file.name.rsplit('.', 1)[0]
                st.session_state['output_filename'] = f"{base_name} edited {time_tag}.docx"
                
                st.success("✅ Processing complete!")

            except Exception as e:
                st.error(f"❌ An error occurred during processing: {str(e)}")
                # Reset state on error so user can retry
                st.session_state['last_processed_id'] = None

        # Display download button if we have processed data for the current file
        if st.session_state.get('last_processed_id') == file_id and 'processed_data' in st.session_state:
            st.download_button(
                label="📥 Download Processed File",
                data=st.session_state['processed_data'],
                file_name=st.session_state['output_filename'],
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key="btn_download"
            )

def calculate_column_widths(data_matrix, start_row_index, usable_width, header_row_idx):
    """
    Calculate optimal column widths based on content length of first 10 rows.
    usable_width is in inches.
    header_row_idx is used to identify specific columns that need more space.
    """
    if not data_matrix:
        return []
    
    num_cols = len(data_matrix[0])
    if num_cols == 0:
        return []
        
    # Content sample: 10 rows starting from start_row_index
    sample_start = min(start_row_index, len(data_matrix))
    sample_end = min(sample_start + 10, len(data_matrix))
    content_sample = data_matrix[sample_start:sample_end]
    
    widths = []
    for col_idx in range(num_cols):
        # Measure max length in the sample
        max_content_len = 0
        if content_sample:
            for row in content_sample:
                val_str = row[col_idx]
                max_content_len = max(max_content_len, len(val_str))
        else:
            max_content_len = 5 # Default fallback
            
        # Heuristic: base width on content length
        # Using a slightly higher multiplier since we ignore headers
        multiplier = 0.8
        min_val = 5
        
        # Specific columns that should not wrap
        if header_row_idx is not None and header_row_idx < len(data_matrix):
            header_text = data_matrix[header_row_idx][col_idx].lower()
            if "country" in header_text or "registration" in header_text:
                multiplier = 2.0 # Give much more weight
                min_val = 15
                
        widths.append(max(min_val, min(max_content_len * multiplier, 50)))
    
    total_raw_width = sum(widths)
    if total_raw_width == 0:
        return [usable_width / num_cols] * num_cols
    
    # Convert to inches and scale
    scaled_widths = [(w / total_raw_width) * usable_width for w in widths]
    
    # Ensure min width (inches)
    for i in range(len(scaled_widths)):
        target_min = 0.5
        if header_row_idx is not None and header_row_idx < len(data_matrix):
            h_text = data_matrix[header_row_idx][i].lower()
            if "country" in h_text:
                target_min = 1.0 # Min 1.0 inch for Country
            elif "registration" in h_text:
                target_min = 1.4 # Min 1.4 inch for Registration Number
        
        if scaled_widths[i] < target_min:
            scaled_widths[i] = target_min
            
    # Re-scale to fit usable width if we went over due to target_min
    current_total = sum(scaled_widths)
    if current_total > usable_width:
        scaled_widths = [(w / current_total) * usable_width for w in scaled_widths]
        
    return scaled_widths

def set_repeat_table_header(row):
    """
    Set a table row to repeat as header on every page.
    """
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), "true")
    trPr.append(tblHeader)

def find_company_name_row(ws):
    """
    Automatically find the row index (1-based) containing "Company Name".
    Searches first 50 rows.
    """
    for row in ws.iter_rows(min_row=1, max_row=50):
        for cell in row:
            if cell.value:
                val_str = str(cell.value).strip().lower()
                if "company name" in val_str:
                    return cell.row
    return 3 # Default fallback if not found

def excel_to_word_landscape(uploaded_file):
    """
    Convert Excel sheet to a landscape Word document with optimized table and merged cells.
    """
    # 1. Read Excel with openpyxl to get merge info
    wb = load_workbook(uploaded_file, data_only=True)
    ws = wb.worksheets[0]
    
    # 2. Automatically detect start row
    header_row = find_company_name_row(ws)
    start_row = header_row + 1
    
    # 2. Get Merged Ranges
    merged_ranges = ws.merged_cells.ranges
    
    # 3. Determine actual data bounds
    max_r = 0
    max_c = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                max_r = max(max_r, cell.row)
                max_c = max(max_c, cell.column)
    
    for m_range in merged_ranges:
        max_r = max(max_r, m_range.max_row)
        max_c = max(max_c, m_range.max_col)
    
    if max_r == 0 or max_c == 0:
        raise ValueError("The Excel file seems to be empty.")

    # 4. Create Word Document
    doc = Document()
    
    # 5. Set Margins and Orientation (A4 Landscape)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    
    section.top_margin = Mm(25.4)
    section.bottom_margin = Mm(22.9)
    section.left_margin = Mm(17.8)
    section.right_margin = Mm(17.8)
    
    usable_width = (section.page_width - section.left_margin - section.right_margin) / 914400
    
    # 6. Create Table
    table = doc.add_table(rows=max_r, cols=max_c)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Set Table to Fixed Layout for stability when copying
    tbl = table._element
    tblPr = tbl.xpath('w:tblPr')
    if not tblPr:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    else:
        tblPr = tblPr[0]
        
    layout = tblPr.xpath('w:tblLayout')
    if not layout:
        layout = OxmlElement('w:tblLayout')
        layout.set(qn('w:val'), 'fixed')
        tblPr.append(layout)
    else:
        layout[0].set(qn('w:val'), 'fixed')
    
    # 7. Extract data matrix efficiently
    data_matrix = []
    # ws.iter_rows with values_only is much faster
    for row in ws.iter_rows(min_row=1, max_row=max_r, min_col=1, max_col=max_c, values_only=True):
        data_matrix.append([str(val) if val is not None else "" for val in row])
    
    # Widths based on 10 rows of content starting from start_row
    col_widths = calculate_column_widths(data_matrix, start_row - 1, usable_width, header_row - 1)
    
    # Pre-identify columns to center
    centered_col_indices = set()
    if header_row - 1 < len(data_matrix):
        # Keywords for columns that should be centered
        center_keywords = [
            "country", "registration", "total net sales are", 
            "consecutive years", "accepted companies", "rejected companies"
        ]
        for i, h_text in enumerate(data_matrix[header_row - 1]):
            h_text_lower = h_text.lower()
            if any(k in h_text_lower for k in center_keywords):
                centered_col_indices.add(i)
    
    # 8. Fill Table and Handle Merges (Headers only)
    header_merged_ranges = [m for m in merged_ranges if m.min_row < start_row]
    
    merged_cells_to_skip = set()
    for m_range in header_merged_ranges:
        for r in range(m_range.min_row, m_range.max_row + 1):
            for c in range(m_range.min_col, m_range.max_col + 1):
                if r == m_range.min_row and c == m_range.min_col:
                    continue
                merged_cells_to_skip.add((r, c))
    
    # Iterate through rows and cells directly (MUCH FASTER than table.cell(r, c))
    for r_idx, word_row in enumerate(table.rows):
        r = r_idx + 1
        for c_idx, word_cell in enumerate(word_row.cells):
            c = c_idx + 1
            
            # Set width (only once per column if we wanted to optimize further, 
            # but once per cell is needed for fixed layout)
            word_cell.width = Inches(col_widths[c_idx])
            
            if (r, c) in merged_cells_to_skip:
                continue
                
            val_str = data_matrix[r_idx][c_idx]
            word_cell.text = val_str
            word_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            
            # Set minimal cell margins (padding) to prevent jumping
            tc = word_cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcMar = OxmlElement('w:tcMar')
            for m_name in ['top', 'left', 'bottom', 'right']:
                node = OxmlElement(f'w:{m_name}')
                node.set(qn('w:w'), '60') # Slightly larger but still compact
                node.set(qn('w:type'), 'dxa')
                tcMar.append(node)
            tcPr.append(tcMar)
            
            # Styling
            is_header = r < start_row
            for para in word_cell.paragraphs:
                if is_header or c_idx in centered_col_indices:
                    para.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # Global Font: Times New Roman, Size 10
                run = para.runs[0] if para.runs else para.add_run()
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                if is_header:
                    run.font.bold = True

    # 9. Execute Merges in Word (Headers only)
    for m_range in header_merged_ranges:
        if m_range.min_row <= max_r and m_range.min_col <= max_c:
            top_left = table.cell(m_range.min_row - 1, m_range.min_col - 1)
            m_max_r = min(m_range.max_row, max_r)
            m_max_c = min(m_range.max_col, max_c)
            bottom_right = table.cell(m_max_r - 1, m_max_c - 1)
            if top_left != bottom_right:
                top_left.merge(bottom_right)

    # 10. Save to BytesIO (NO REPEATING HEADER)
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output

def render_matrix_processor_tab():
    st.markdown("<h2 style='text-align: center; color: #1e3a8a;'>Matrix Processor</h2>", unsafe_allow_html=True)
    st.markdown("<p class='subtitle'>Convert Excel matrices to professional landscape Word documents</p>", unsafe_allow_html=True)

    # Settings
    auto_process = st.checkbox("Auto-process after upload", value=True, key="matrix_auto_process_check")

    # File Uploader
    uploaded_file = st.file_uploader(
        "Upload Excel file (.xlsx)",
        type=["xlsx"],
        help="Reads the first sheet of the Excel file",
        key="matrix_uploader"
    )

    if uploaded_file is not None:
        file_id = f"matrix_{uploaded_file.name}_{uploaded_file.size}"
        
        # Process button
        start_manual = st.button(
            "🚀 Convert to Word", 
            key="btn_matrix_process", 
            use_container_width=True,
            disabled=auto_process
        )
        
        # Trigger processing if manual button clicked OR (auto-process is ON and file changed)
        should_run = start_manual or (auto_process and st.session_state.get('last_matrix_id') != file_id)

        if should_run:
            st.session_state['last_matrix_id'] = file_id
            
            with st.status("Processing Excel to Word...", expanded=True) as status:
                try:
                    status.write("Reading Excel data...")
                    # Reset pointer just in case
                    uploaded_file.seek(0)
                    
                    status.write("Generating Word document (Landscape)...")
                    processed_stream = excel_to_word_landscape(uploaded_file)
                    
                    status.write("Finalizing...")
                    st.session_state['matrix_processed_data'] = processed_stream.getvalue()
                    
                    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
                    base_name = uploaded_file.name.rsplit('.', 1)[0]
                    st.session_state['matrix_output_filename'] = f"{base_name}_{timestamp}.docx"
                    
                    status.update(label="✅ Conversion complete!", state="complete", expanded=False)
                    st.success("Successfully converted to Word!")
                except Exception as e:
                    st.error(f"❌ Error during conversion: {str(e)}")
                    st.session_state['last_matrix_id'] = None

        # Download button
        if st.session_state.get('last_matrix_id') == file_id and 'matrix_processed_data' in st.session_state:
            st.download_button(
                label="📥 Download Word Document",
                data=st.session_state['matrix_processed_data'],
                file_name=st.session_state['matrix_output_filename'],
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key="btn_matrix_download"
            )

def main():
    # Application Title
    st.markdown("<h1 style='text-align: center; color: #1e3a8a;'>Ultimate Tool Hub</h1>", unsafe_allow_html=True)
    
    # Define Tabs
    tab_titles = ["📄 Comparable Taxpayers Processor", "📊 Matrix Processor", "🛠️ Other Tools"]
    tabs = st.tabs(tab_titles)

    with tabs[0]:
        word_processor_tab()

    with tabs[1]:
        render_matrix_processor_tab()

    with tabs[2]:
        st.markdown("<h2 style='text-align: center; color: #1e3a8a;'>Coming Soon</h2>", unsafe_allow_html=True)
        st.info("We are developing new tools to make your work more efficient. Stay tuned!")

    # Footer
    st.markdown("---")
    st.markdown(
        "<div style='text-align: center; color: #94a3b8; font-size: 0.8rem;'>"
        "Efficiency Optimization System."
        "</div>", 
        unsafe_allow_html=True
    )

if __name__ == "__main__":
    main()
