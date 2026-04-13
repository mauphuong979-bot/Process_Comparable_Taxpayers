from __future__ import annotations

import re
import io
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

STYLES_TO_NORMALIZE = {
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Strong",
    "Section_Level1",
    "Section_Level2",
    "Section_Level3",
    "Section_Level4",
    "Section_Level4App",
    "Section_Level5",
}

STYLE_FONT_SIZE_MAP = {
    "Heading 1": 15,
    "Section_Level1": 15,
    "Heading 2": 13,
    "Section_Level2": 13,
    "Heading 3": 12,
    "Section_Level3": 12,
    "Section_Level4": 12,
}

PAGE_BREAK_KEYWORDS = [
    "Comparable Taxpayers",
    "Step One: Commercial Database Search",
    "Step Two: Bulk Rejection of Uncontrolled Taxpayers",
    "Step Three: Qualitative Review Rejection of Uncontrolled Taxpayers",
    "Final Set of Comparable Taxpayers",
    "Arm's Length Range",
    "APPENDIX A",
    "APPENDIX B",
]

SPECIAL_TARGET_TEXT = "Unconsolidated Income Statement"
DEFAULT_FONT_NAME = "Times New Roman"

# =========================
# HÀM TIỆN ÍCH
# =========================

def normalize_compare_text(text: str) -> str:
    """
    Chuẩn hóa text để so sánh.
    """
    if text is None:
        return ""
    # Loại bỏ ký tự đặc biệt hay gặp trong Word
    text = text.replace("\r", "").replace("\n", "").replace("\t", "")
    text = text.replace("\u00a0", " ")  # non-breaking space
    text = text.strip()
    text = text.replace("\u2019", "'")  # smart quote
    text = text.replace("\u2018", "'")  # smart quote open
    return text

def remove_outline_level(paragraph) -> None:
    """
    Xóa thẻ w:outlineLvl trong XML của paragraph pPr.
    """
    pPr = paragraph._element.get_or_add_pPr()
    outlines = pPr.xpath('./w:outlineLvl')
    for ol in outlines:
        pPr.remove(ol)

def is_in_table(paragraph) -> bool:
    """
    Kiểm tra paragraph có nằm trong table không bằng cách tìm tổ tiên là thẻ tc (table cell).
    """
    return len(paragraph._p.xpath('./ancestor::w:tc')) > 0

# =========================
# CÁC BƯỚC XỬ LÝ
# =========================

def normalize_heading_styles(doc: Document) -> None:
    """
    Đổi các style heading/strong/custom về Normal.
    Giữ màu đen, chữ đậm, và set size theo style cũ nếu có.
    """
    for para in doc.paragraphs:
        style_name = para.style.name
        
        # Kiểm tra xem paragraph có cần chuẩn hóa không:
        # 1. Style của paragraph nằm trong list chuẩn hóa
        # 2. Hoặc paragraph có chứa run mang style "Strong"
        has_strong_run = any(run.style and run.style.name == "Strong" for run in para.runs)
        
        if style_name in STYLES_TO_NORMALIZE or has_strong_run:
            # Lưu lại size trước khi đổi style nếu có trong map
            new_size = STYLE_FONT_SIZE_MAP.get(style_name)
            
            para.style = doc.styles['Normal']
            
            # Xóa Outline Level bằng can thiệp XML
            remove_outline_level(para)
            
            # Áp dụng định dạng cho tất cả các run trong paragraph
            for run in para.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0)
                if new_size:
                    run.font.size = Pt(new_size)

def remove_leading_spaces_tabs(doc: Document) -> None:
    """
    Xóa khoảng trắng/tab ở đầu từng paragraph, bảo toàn định dạng.
    """
    for para in doc.paragraphs:
        if not para.runs:
            continue
            
        # Ta duyệt qua các run cho đến khi tìm thấy text
        for run in para.runs:
            if run.text:
                original_text = run.text
                new_text = original_text.lstrip(' \t')
                run.text = new_text
                
                # Nếu run này có chứa text và sau khi lstrip vẫn còn hoặc đã hết text
                # thì ta dừng việc xóa ở đầu paragraph này.
                if new_text or original_text:
                    # Nếu run này trắng hoàn toàn sau khi lstrip, ta tiếp tục sang run sau
                    if not new_text:
                        continue
                    else:
                        break

def replace_manual_page_breaks(doc: Document) -> None:
    """
    Thay manual page break bằng paragraph mark (trong python-docx là xóa break element).
    """
    for para in doc.paragraphs:
        for run in para.runs:
            # Tìm thẻ <w:br w:type="page"/> trong XML của run
            brs = run._element.xpath('.//w:br[@w:type="page"]')
            for br in brs:
                br.getparent().remove(br)

def remove_extra_blank_lines_outside_tables(doc: Document) -> None:
    """
    Xóa dòng trống thừa ngoài table, chỉ giữ tối đa 1 dòng trống liên tiếp.
    """
    paragraphs_to_remove = []
    prev_blank = False
    
    # Do python-docx không hỗ trợ xóa trực tiếp dễ dàng khi đang loop,
    # ta sẽ đánh dấu các paragraph cần xóa.
    for para in doc.paragraphs:
        if is_in_table(para):
            prev_blank = False
            continue
            
        text = para.text.strip()
        if not text:
            if prev_blank:
                paragraphs_to_remove.append(para)
            else:
                prev_blank = True
        else:
            prev_blank = False
            
    # Xóa các paragraph đã đánh dấu
    for para in paragraphs_to_remove:
        p = para._element
        p.getparent().remove(p)
        para._p = para._element = None

def insert_page_break_before_keywords(doc: Document, keywords: Iterable[str]) -> None:
    """
    Chèn page break trước paragraph nếu text khớp keyword.
    """
    keyword_set = {normalize_compare_text(k) for k in keywords}
    
    # Ta cần lặp qua một bản sao của list paragraphs vì ta sẽ chèn thêm paragraph mới
    for i, para in enumerate(list(doc.paragraphs)):
        if is_in_table(para):
            continue
            
        norm_para_text = normalize_compare_text(para.text)
        # Sử dụng so sánh linh hoạt hơn: bắt đầu bằng keyword hoặc khớp hoàn toàn
        is_match = False
        for kw in keyword_set:
            if norm_para_text == kw or norm_para_text.startswith(kw + ":") or norm_para_text.startswith(kw + " "):
                is_match = True
                break
        
        # CHỈ chèn ngắt trang nếu không phải là đoạn đầu tiên của tài liệu
        if is_match and i > 0:
            new_p = para.insert_paragraph_before()
            new_p.add_run().add_break(WD_BREAK.PAGE)

def insert_special_break_before_income_statement(doc: Document, target_text: str) -> None:
    """
    Chèn page break ở trước paragraph nằm trên target 2 đoạn.
    """
    target_norm = normalize_compare_text(target_text)
    paragraphs = list(doc.paragraphs)
    
    for i, para in enumerate(paragraphs):
        if is_in_table(para):
            continue
            
        if normalize_compare_text(para.text) == target_norm:
            target_idx = i - 2
            if target_idx >= 0:
                # Chèn trước paragraph tại target_idx
                paragraphs[target_idx].insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)

def set_document_font(doc: Document, font_name: str = DEFAULT_FONT_NAME) -> None:
    """
    Đặt font "Times New Roman" cho toàn bộ tài liệu bao gồm body, tables, headers và footers.
    Sử dụng kỹ thuật can thiệp XML để đảm bảo font được áp dụng triệt để.
    """
    # Helper để áp font triệt để cho run
    def apply_run_font(run):
        run.font.name = font_name
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:eastAsia'), font_name)
        rFonts.set(qn('w:cs'), font_name)

    # 1. Chỉnh font cho tất cả Styles có trong document
    for style in doc.styles:
        if hasattr(style, 'font'):
            try:
                style.font.name = font_name
                # Can thiệp XML cho style font
                rPr = style.element.get_or_add_rPr()
                rFonts = rPr.get_or_add_rFonts()
                rFonts.set(qn('w:ascii'), font_name)
                rFonts.set(qn('w:hAnsi'), font_name)
                rFonts.set(qn('w:eastAsia'), font_name)
                rFonts.set(qn('w:cs'), font_name)
            except Exception:
                continue

    # 2. Helper xử lý Table đệ quy (để xử lý cả nested tables)
    def process_table_recursive(table):
        for row in table.rows:
            for cell in row.cells:
                # Xử lý các paragraph trong cell
                for para in cell.paragraphs:
                    for run in para.runs:
                        apply_run_font(run)
                # Đệ quy nếu cell chứa table con
                for nested_table in cell.tables:
                    process_table_recursive(nested_table)

    # 3. Xử lý nội dung chính (body)
    for para in doc.paragraphs:
        for run in para.runs:
            apply_run_font(run)
    
    for table in doc.tables:
        process_table_recursive(table)

    # 4. Xử lý Headers và Footers của tất cả các Section
    for section in doc.sections:
        # Header
        header = section.header
        for para in header.paragraphs:
            for run in para.runs:
                apply_run_font(run)
        for table in header.tables:
            process_table_recursive(table)
            
        # Footer
        footer = section.footer
        for para in footer.paragraphs:
            for run in para.runs:
                apply_run_font(run)
        for table in footer.tables:
            process_table_recursive(table)

def process_docx(file_stream: io.BytesIO) -> io.BytesIO:
    """
    Hàm chính xử lý Document và trả về stream mới.
    """
    doc = Document(file_stream)
    
    normalize_heading_styles(doc)
    remove_leading_spaces_tabs(doc)
    replace_manual_page_breaks(doc)
    remove_extra_blank_lines_outside_tables(doc)
    insert_page_break_before_keywords(doc, PAGE_BREAK_KEYWORDS)
    insert_special_break_before_income_statement(doc, SPECIAL_TARGET_TEXT)
    set_document_font(doc, DEFAULT_FONT_NAME)
    
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
