from docx import Document
import os

def inspect_docx(file_path):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    doc = Document(file_path)
    
    print(f"--- Document: {file_path} ---")
    print(f"Number of paragraphs: {len(doc.paragraphs)}")
    print(f"Number of tables: {len(doc.tables)}")
    
    print("\n--- Styles defined in document ---")
    styles = {s.name for s in doc.styles}
    print(sorted(list(styles)))
    
    print("\n--- Sample Paragraphs (First 20) ---")
    for i, para in enumerate(doc.paragraphs[:20]):
        print(f"{i}: Style: [{para.style.name}] | Text: {para.text[:100]!r}")
        
    print("\n--- Searching for keywords ---")
    keywords = [
        "Step One: Commercial Database Search",
        "Step Two: Bulk Rejection of Uncontrolled Taxpayers",
        "Step Three: Qualitative Review Rejection of Uncontrolled Taxpayers",
        "Final Set of Comparable Taxpayers",
        "Arm's Length Range",
        "APPENDIX A",
        "APPENDIX B",
        "Unconsolidated Income Statement",
        "Consolidated Income Statement"
    ]
    
    word_found = {k: False for k in keywords}
    for para in doc.paragraphs:
        txt = para.text.strip()
        for k in keywords:
            if k in txt:
                word_found[k] = True
                
    for k, found in word_found.items():
        print(f"Keyword '{k}': {'FOUND' if found else 'NOT FOUND'}")

    print("\n--- Table structures ---")
    for i, table in enumerate(doc.tables[:3]): # Check first 3 tables
        print(f"Table {i}: {len(table.rows)} rows, {len(table.columns)} columns")
        # Check first cell of first row
        if len(table.rows) > 0 and len(table.columns) > 0:
            print(f"  First cell text: {table.cell(0, 0).text[:50]!r}")

if __name__ == "__main__":
    inspect_docx(r"d:\AI\Python\GitHub\Fix Comparable Taxpayers\Demo_2025 avg Report.docx")
