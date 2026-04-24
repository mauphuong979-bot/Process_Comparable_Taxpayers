import io
from docx import Document
from processor import normalize_compare_text, SPECIAL_TARGET_TEXTS

def test_logic():
    doc = Document()
    doc.add_paragraph("Para 1")
    doc.add_paragraph("Para 2")
    doc.add_paragraph("Consolidated  Income Statement") # Double space
    doc.add_paragraph("Para 4")
    doc.add_paragraph("CONSOLIDATED INCOME STATEMENT") # All caps
    doc.add_paragraph("Para 6")
    doc.add_paragraph("Consolidated\tIncome\tStatement") # Tabs
    doc.add_paragraph("Para 8")
    doc.add_paragraph("Consolidated Income Statement ") # Trailing space
    
    target_norms = {normalize_compare_text(t) for t in SPECIAL_TARGET_TEXTS}
    print(f"Target norms: {target_norms}")
    
    found = False
    for i, para in enumerate(doc.paragraphs):
        norm_text = normalize_compare_text(para.text)
        print(f"Para {i}: {para.text!r} -> {norm_text!r}")
        if norm_text in target_norms:
            print(f"Found match at index {i}")
            target_idx = i - 2
            if target_idx >= 0:
                print(f"Inserting break before index {target_idx} (text: {doc.paragraphs[target_idx].text!r})")
                doc.paragraphs[target_idx].insert_paragraph_before().add_run().add_break()
                found = True
    
    if not found:
        print("Match not found!")

if __name__ == "__main__":
    test_logic()
