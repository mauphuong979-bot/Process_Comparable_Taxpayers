from docx import Document
import re

def normalize_compare_text(text: str) -> str:
    if text is None:
        return ""
    text = text.replace("\r", "").replace("\n", "").replace("\t", "")
    text = text.strip()
    text = text.replace("\u2019", "'")
    return text

def inspect_docx(file_path):
    doc = Document(file_path)
    
    print("--- Checking for Arm's Length Range ---")
    for i, para in enumerate(doc.paragraphs):
        norm_txt = normalize_compare_text(para.text)
        if "Arm" in norm_txt or "Length" in norm_txt or "Range" in norm_txt:
            print(f"Para {i}: Original: {para.text!r} | Normalized: {norm_txt!r}")

    print("\n--- Checking Styles of Headings ---")
    for i, para in enumerate(doc.paragraphs):
        if i > 50: break # sample
        if para.text.strip():
            print(f"Para {i}: Style: {para.style.name} | Text: {para.text[:50]!r}")

    print("\n--- Checking for Section_Level styles ---")
    section_paras = [p for p in doc.paragraphs if "Section_Level" in p.style.name]
    for p in section_paras[:10]:
        print(f"Style: {p.style.name} | Text: {p.text!r}")

if __name__ == "__main__":
    inspect_docx(r"d:\AI\Python\GitHub\Fix Comparable Taxpayers\Demo_2025 avg Report.docx")
