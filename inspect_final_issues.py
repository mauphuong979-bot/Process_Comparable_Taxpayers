from docx import Document

def inspect_font_and_structure(file_path):
    doc = Document(file_path)
    
    # 1. Check first paragraph
    if doc.paragraphs:
        first_para = doc.paragraphs[0]
        print(f"First Para Text: {first_para.text!r}")
        print(f"First Para Style: {first_para.style.name}")

    # 2. Check fonts in tables
    print("\n--- Checking Fonts in Tables ---")
    for i, table in enumerate(doc.tables[:3]): # Check first 3 tables
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.font.name:
                            print(f"Table {i} Run Font: {run.font.name} | Text: {run.text[:20]!r}")
    
    # 3. Check fonts in headers/footers
    print("\n--- Checking Fonts in Headers/Footers ---")
    for i, section in enumerate(doc.sections):
        header = section.header
        for para in header.paragraphs:
            for run in para.runs:
                if run.font.name:
                    print(f"Section {i} Header Font: {run.font.name}")
        
        footer = section.footer
        for para in footer.paragraphs:
            for run in para.runs:
                if run.font.name:
                    print(f"Section {i} Footer Font: {run.font.name}")

if __name__ == "__main__":
    inspect_font_and_structure(r"d:\AI\Python\GitHub\Fix Comparable Taxpayers\Demo_2025 avg Report.docx")
