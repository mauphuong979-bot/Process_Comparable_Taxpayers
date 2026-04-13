import io
from processor import process_docx

def test_processor():
    # Use the existing demo file if it exists, otherwise create a minimal one
    input_path = "Demo_2025 avg Report.docx"
    try:
        with open(input_path, "rb") as f:
            content = f.read()
            print(f"Read {len(content)} bytes from {input_path}")
    except FileNotFoundError:
        print("Demo file not found, creating a test file...")
        from docx import Document
        doc = Document()
        doc.add_heading("Heading 1 Test", level=1)
        doc.add_paragraph("  Leading spaces test")
        doc.add_paragraph("Step One: Commercial Database Search")
        doc.add_paragraph("Normal text")
        doc.add_paragraph("")
        doc.add_paragraph("")
        doc.add_paragraph("Another normal text")
        doc.add_paragraph("Unconsolidated Income Statement")
        
        mem = io.BytesIO()
        doc.save(mem)
        content = mem.getvalue()
        print("Created test content in memory.")

    input_stream = io.BytesIO(content)
    try:
        output_stream = process_docx(input_stream)
        with open("test_output.docx", "wb") as f:
            f.write(output_stream.getbuffer())
        print("Successfully processed and saved to test_output.docx")
    except Exception as e:
        print(f"Error during processing: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    test_processor()
