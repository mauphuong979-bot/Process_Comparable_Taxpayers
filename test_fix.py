from docx import Document
from docx.enum.section import WD_SECTION
import io

def test_normalization():
    path = r"d:\AI\Python\GitHub\Fix Comparable Taxpayers\Demo_2025 avg Report.docx"
    doc = Document(path)
    
    # Target Para 36: Bulk Rejections Summary
    target_para = None
    for para in doc.paragraphs:
        if "Bulk Rejections Summary" in para.text:
            target_para = para
            break
            
    if not target_para:
        print("Could not find target paragraph")
        return

    print(f"BEFORE: Style: {target_para.style.name}, Outline: {target_para.paragraph_format.outline_level}")
    
    # Apply changes
    target_para.style = doc.styles['Normal']
    target_para.paragraph_format.outline_level = None
    
    # Save to a temporary stream and reload
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    
    doc2 = Document(stream)
    target_para2 = None
    for para in doc2.paragraphs:
        if "Bulk Rejections Summary" in para.text:
            target_para2 = para
            break
            
    print(f"AFTER: Style: {target_para2.style.name}, Outline: {target_para2.paragraph_format.outline_level}")

if __name__ == "__main__":
    test_normalization()
